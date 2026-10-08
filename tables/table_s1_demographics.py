#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Table S1 — demographic characteristics of the FDSC, CFPS 2012 and MIDUS M2 cohorts.

For each cohort the low-symptom group (CES-D total < 16) and the high-symptom
group (CES-D total >= 16) are compared on sex and age.

Conventions (identical to the RCT tables so that all demographic tables of the
manuscript share one rule):

* sex — uncorrected Pearson chi-squared on the 2x2 low/high x male/female
  table; Fisher's exact test is substituted when the smallest *expected*
  frequency falls below 5.
* age — independent-samples t test on low vs high (low - high).  Levene's test
  decides the variance assumption: pooled (Student) when P >= 0.05, Welch when
  P < 0.05.  Only MIDUS M2 requires Welch (Levene P = 0.031, t = 5.16); FDSC
  (t = -0.51, P = 0.611) and CFPS 2012 (t = -15.45, P < 0.001) use the pooled
  test.  This reproduces Table S1 cell for cell.

Inputs (data/, see data/README.md):
    3447例-CESD数据已反向计分(3).xlsx    FDSC,  columns 性别 (1 = male, 2 = female),
                                          年龄, CESD_1 .. CESD_20
    CFPS_cesd20_scores.csv                CFPS 2012, columns gender_name, age,
                                          cesd20_total
    MIDUS_CESD20.xlsx                     MIDUS M2, columns Sex_label, Age,
                                          CESD20_total

Outputs (results/):
    table_s1_demographics.csv
    table_s1_demographics.md
    表S1_三队列人口学特征.docx

Usage:
    python tables/table_s1_demographics.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

_HERE = Path(__file__).resolve().parent
_REPO = _HERE.parent
for _p in (str(_REPO), str(_HERE)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from config import DATA_DIR, OUT_DIR  # noqa: E402

# CES-D cut-off separating the low- from the high-symptom group.
CESD_CUTOFF = 16
CSV = OUT_DIR / "table_s1_demographics.csv"
MD = OUT_DIR / "table_s1_demographics.md"
DOCX = OUT_DIR / "表S1_三队列人口学特征.docx"


def load_cohorts():
    """Return an ordered list of (label, sex-is-male flag, age, CES-D total)."""
    fdsc = pd.read_excel(DATA_DIR / "3447例-CESD数据已反向计分(3).xlsx")
    fdsc_total = fdsc[[f"CESD_{i}" for i in range(1, 21)]].apply(
        pd.to_numeric, errors="coerce").sum(axis=1)
    cohorts = [("FDSC",
                fdsc["性别"].astype(float) == 1,
                pd.to_numeric(fdsc["年龄"], errors="coerce").astype(float),
                fdsc_total.astype(float))]

    cfps = pd.read_csv(DATA_DIR / "CFPS_cesd20_scores.csv", encoding="utf-8-sig")
    cohorts.append(("CFPS 2012",
                    cfps["gender_name"].astype(str).str.strip() == "男",
                    pd.to_numeric(cfps["age"], errors="coerce").astype(float),
                    pd.to_numeric(cfps["cesd20_total"], errors="coerce").astype(float)))

    midus = pd.read_excel(DATA_DIR / "MIDUS_CESD20.xlsx")
    cohorts.append(("MIDUS M2",
                    midus["Sex_label"].astype(str).str.strip() == "Male",
                    pd.to_numeric(midus["Age"], errors="coerce").astype(float),
                    pd.to_numeric(midus["CESD20_total"], errors="coerce").astype(float)))
    return cohorts


def _fmt_p(p: float) -> str:
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def _n_pct(k: int, n: int) -> str:
    return f"{k:,}({100 * k / n:.1f}%)"


def _mean_sd(x: np.ndarray) -> str:
    x = x[~np.isnan(x)]
    return f"{np.mean(x):.2f}±{np.std(x, ddof=1):.2f}"


def _sex_test(male: pd.Series, group: pd.Series) -> tuple[str, str]:
    """Sex comparison between two groups; returns (statistic, P)."""
    table = np.array([[int((male & group).sum()), int((~male & group).sum())],
                      [int((male & ~group).sum()), int((~male & ~group).sum())]])
    expected = (table.sum(axis=1, keepdims=True) * table.sum(axis=0, keepdims=True)
                / table.sum())
    if expected.min() < 5:
        return "—", _fmt_p(stats.fisher_exact(table)[1])
    chi2, p = stats.chi2_contingency(table, correction=False)[:2]
    return f"{chi2:.2f}", _fmt_p(p)


def _age_test(low_age: np.ndarray, high_age: np.ndarray) -> tuple[str, str]:
    """Age comparison between two groups; pooled t unless Levene rejects."""
    a = low_age[~np.isnan(low_age)]
    b = high_age[~np.isnan(high_age)]
    equal_var = stats.levene(a, b).pvalue >= 0.05
    t, p = stats.ttest_ind(a, b, equal_var=equal_var)
    return f"{t:.2f}", _fmt_p(p)


def build_table(cohorts) -> pd.DataFrame:
    rows = []
    for label, male, age, total in cohorts:
        low = total < CESD_CUTOFF
        high = total >= CESD_CUTOFF
        n, n_lo, n_hi = len(total), int(low.sum()), int(high.sum())
        n_male = int(male.sum())
        sex_stat, sex_p = _sex_test(male, low)
        age_stat, age_p = _age_test(age[low].values, age[high].values)
        rows.append({"Cohort": label, "Variable": "Sex",
                     "Total sample": "", "Low-symptom group": "",
                     "High-symptom group": "", "χ²/t": "", "P": ""})
        rows.append({"Cohort": "", "Variable": "Male, n (%)",
                     "Total sample": _n_pct(n_male, n),
                     "Low-symptom group": _n_pct(int((male & low).sum()), n_lo),
                     "High-symptom group": _n_pct(int((male & high).sum()), n_hi),
                     "χ²/t": sex_stat, "P": sex_p})
        rows.append({"Cohort": "", "Variable": "Female, n (%)",
                     "Total sample": _n_pct(n - n_male, n),
                     "Low-symptom group": _n_pct(int((~male & low).sum()), n_lo),
                     "High-symptom group": _n_pct(int((~male & high).sum()), n_hi),
                     "χ²/t": "", "P": ""})
        rows.append({"Cohort": "", "Variable": "Age (years)",
                     "Total sample": _mean_sd(age.values),
                     "Low-symptom group": _mean_sd(age[low].values),
                     "High-symptom group": _mean_sd(age[high].values),
                     "χ²/t": age_stat, "P": age_p})
        print(f"[{label}] N = {n:,} (low {n_lo:,} / high {n_hi:,}) | "
              f"sex chi2 = {sex_stat}, P = {sex_p} | age t = {age_stat}, P = {age_p}")
    return pd.DataFrame(rows)


def to_markdown(df: pd.DataFrame) -> str:
    head = list(df.columns)
    lines = ["| " + " | ".join(head) + " |",
             "|" + "|".join(["---"] * len(head)) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(str(r[c]) for c in head) + " |")
    return "\n".join(lines) + "\n"


def to_docx(df: pd.DataFrame, path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ImportError:
        print("[skip] python-docx not installed; docx not written")
        return

    doc = Document()
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run("Table S1 Demographic characteristics of the FDSC, CFPS 2012 "
                      "and MIDUS M2 cohorts")
    run.bold = True
    run.font.size = Pt(10.5)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(attr), "Times New Roman" if attr != "w:eastAsia" else "宋体")

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, h in enumerate(df.columns):
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(row[h]))
            r.font.size = Pt(10)
    doc.save(path)
    print(f"[export] {path}")


def main() -> None:
    df = build_table(load_cohorts())
    df.to_csv(CSV, index=False, encoding="utf-8-sig")
    MD.write_text(to_markdown(df), encoding="utf-8")
    print(f"[export] {CSV}\n[export] {MD}")
    to_docx(df, DOCX)


if __name__ == "__main__":
    main()
